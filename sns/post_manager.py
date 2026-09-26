"""X投稿候補の管理（表計算 sns/x_posts.xlsx）。

毎日の流れ:
    1. add-sources  day.json   分析した伸びている投稿（分析元）を記録
    2. add-types    day.json   導き出した型を記録
    3. add-candidates day.json 候補を追加。既存の候補・採用済み・戦略書テンプレと
                               似ているものは追加せず削除ログへ
    4. present 日付 ID×5       その日の提示（1〜5番）を記録
    5. adopt 日付 番号...      採用された番号にチェック。採用済みは二度と提示しない
    status                     件数の確認

JSON はいずれも {"date": "YYYY-MM-DD", "items": [...]} の形。
"""

import json
import re
import sys
import unicodedata
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
from x_count import weighted_length  # noqa: E402

BOOK = Path(__file__).parent / "x_posts.xlsx"
STRATEGY = Path(__file__).parent / "x-strategy.md"
SIMILAR = 0.5  # これ以上似ていたら削除

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
ADOPT_FILL = PatternFill("solid", fgColor="E2EFDA")
SHOWN_FILL = PatternFill("solid", fgColor="FFF2CC")

SHEETS = {
    "候補": ["ID", "作成日", "型", "本文", "X換算(/280)", "行数", "1行目字数",
             "状態", "提示日", "提示番号", "採用", "採用日", "最も似ている既存", "類似度"],
    "型": ["日付", "型ID", "型名", "特徴", "分析元の件数", "代表例"],
    "分析元": ["日付", "No", "アカウント", "URL", "本文（抜粋）", "型ID", "反応数"],
    "削除ログ": ["日付", "削除した本文", "似ていた相手", "類似度"],
}
WIDTHS = {
    "候補": [7, 11, 16, 48, 10, 6, 9, 8, 11, 8, 6, 11, 16, 7],
    "型": [11, 6, 22, 50, 10, 40],
    "分析元": [11, 5, 18, 40, 50, 6, 10],
    "削除ログ": [11, 50, 18, 7],
}
COL = {name: i + 1 for i, name in enumerate(SHEETS["候補"])}


def style_header(ws, widths):
    for i, cell in enumerate(ws[1]):
        cell.font = Font(name=FONT, bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = widths[i]
    ws.freeze_panes = "A2"


def new_book():
    wb = Workbook()
    guide = wb.active
    guide.title = "使い方"
    for row in [
        ["X投稿候補の管理表"],
        [],
        ["シート", "内容"],
        ["候補", "投稿候補。『採用』列に ✅ が付いたものは使用済みで、二度と提示しない"],
        ["型", "毎日の分析で導き出した『伸びる型』"],
        ["分析元", "分析した伸びている投稿（URL付き）"],
        ["削除ログ", "既存の候補・採用済み・戦略書のテンプレと似ていたため削除した候補"],
        [],
        ["状態の意味", ""],
        ["候補", "まだ提示していない"],
        ["提示中", "その日の5案として提示した（未採用なら後日また提示されることがある）"],
        ["採用", "採用済み。再提示しない"],
        [],
        ["類似判定", f"空白・記号を除いた本文で、文字の2文字組の重なりか並びの一致率が {SIMILAR} 以上なら『似ている』"],
        ["反応数", "ウェブ検索経由で集めた投稿は、いいね数などが取れないため『不明』"],
    ]:
        guide.append(row)
    guide["A1"].font = Font(name=FONT, bold=True, size=14)
    for r in (3, 9):
        guide.cell(r, 1).font = Font(name=FONT, bold=True)
    guide.column_dimensions["A"].width = 14
    guide.column_dimensions["B"].width = 90
    for name, head in SHEETS.items():
        ws = wb.create_sheet(name)
        ws.append(head)
        style_header(ws, WIDTHS[name])
    dv = DataValidation(type="list", formula1='"✅"', allow_blank=True)
    wb["候補"].add_data_validation(dv)
    dv.add(f"K2:K5000")
    return wb


def open_book():
    return load_workbook(BOOK) if BOOK.exists() else new_book()


def normalize(text):
    text = unicodedata.normalize("NFKC", text)
    return "".join(ch for ch in text if unicodedata.category(ch)[0] in "LN")


def similarity(a, b):
    a, b = normalize(a), normalize(b)
    if not a or not b:
        return 0.0
    ga = {a[i:i + 2] for i in range(len(a) - 1)}
    gb = {b[i:i + 2] for i in range(len(b) - 1)}
    jaccard = len(ga & gb) / len(ga | gb) if ga | gb else 0.0
    return max(jaccard, SequenceMatcher(None, a, b).ratio())


def strategy_templates():
    """戦略書の5章（テンプレ集）のコードブロックを比較対象にする。"""
    if not STRATEGY.exists():
        return []
    md = STRATEGY.read_text(encoding="utf-8")
    sec = md.split("## 5.", 1)[-1].split("## 6.", 1)[0]
    blocks = re.findall(r"\*\*([^*\n]+)\*\*\n```\n(.*?)\n```", sec, flags=re.S)
    return [(f"戦略書テンプレ{label.split()[0]}", body) for label, body in blocks]


def load_json(path):
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    return data["date"], data["items"]


def write_row(ws, values, fill=None):
    ws.append(values)
    for cell in ws[ws.max_row]:
        cell.font = Font(name=FONT)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if fill:
            cell.fill = fill


def add_sources(path):
    date, items = load_json(path)
    wb = open_book()
    ws = wb["分析元"]
    for i, it in enumerate(items, 1):
        write_row(ws, [date, i, it["account"], it["url"], it["text"], it["type"], it.get("metrics", "不明")])
    wb.save(BOOK)
    print(f"分析元 {len(items)} 件を記録")


def add_types(path):
    date, items = load_json(path)
    wb = open_book()
    ws = wb["型"]
    for it in items:
        write_row(ws, [date, it["id"], it["name"], it["feature"], it["count"], it["example"]])
    wb.save(BOOK)
    print(f"型 {len(items)} 件を記録")


def next_id(ws):
    nums = [int(str(r[0])[1:]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    return max(nums, default=0) + 1


def add_candidates(path):
    date, items = load_json(path)
    wb = open_book()
    ws, log = wb["候補"], wb["削除ログ"]
    pool = [(r[0], r[3]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    pool += strategy_templates()
    n = next_id(ws)
    added = dropped = 0
    for it in items:
        body = it["text"].strip("\n")
        best_id, best = None, 0.0
        for pid, ptext in pool:
            s = similarity(body, ptext)
            if s > best:
                best_id, best = pid, s
        if best >= SIMILAR:
            write_row(log, [date, body, best_id, round(best, 2)])
            print(f"削除: {body.splitlines()[0]} … ↔ {best_id}（{best:.2f}）")
            dropped += 1
            continue
        cid = f"P{n:04d}"
        n += 1
        lines = body.split("\n")
        write_row(ws, [cid, date, it["type"], body, weighted_length(body), len(lines),
                       len(lines[0]), "候補", None, None, None, None, best_id, round(best, 2)])
        pool.append((cid, body))
        added += 1
    wb.save(BOOK)
    print(f"候補 {added} 件を追加 / 似ていたため {dropped} 件を削除")


def rows(ws):
    return list(ws.iter_rows(min_row=2))


def present(date, ids):
    wb = open_book()
    ws = wb["候補"]
    by_id = {r[0].value: r for r in rows(ws)}
    for no, cid in enumerate(ids, 1):
        r = by_id[cid]
        if r[COL["採用"] - 1].value:
            sys.exit(f"{cid} は採用済みなので提示できません")
    for no, cid in enumerate(ids, 1):
        r = by_id[cid]
        r[COL["状態"] - 1].value = "提示中"
        r[COL["提示日"] - 1].value = date
        r[COL["提示番号"] - 1].value = no
        for c in r:
            c.fill = SHOWN_FILL
        print(f"{no}. {cid}")
    wb.save(BOOK)


def adopt(date, nums):
    wb = open_book()
    ws = wb["候補"]
    hit = []
    for r in rows(ws):
        if r[COL["提示日"] - 1].value == date and r[COL["提示番号"] - 1].value in nums:
            r[COL["状態"] - 1].value = "採用"
            r[COL["採用"] - 1].value = "✅"
            r[COL["採用日"] - 1].value = date
            for c in r:
                c.fill = ADOPT_FILL
            hit.append((r[COL["提示番号"] - 1].value, r[0].value))
    missing = set(nums) - {h[0] for h in hit}
    if missing:
        sys.exit(f"{date} に提示した番号 {sorted(missing)} が見つかりません")
    # 同じ日に提示して採用されなかったものは、また提示できるよう候補に戻す
    for r in rows(ws):
        if r[COL["提示日"] - 1].value == date and r[COL["状態"] - 1].value == "提示中":
            r[COL["状態"] - 1].value = "候補"
            for c in r:
                c.fill = PatternFill(fill_type=None)
    wb.save(BOOK)
    for no, cid in sorted(hit):
        print(f"✅ {no}. {cid}")


def status():
    wb = open_book()
    counts = {}
    for r in wb["候補"].iter_rows(min_row=2, values_only=True):
        counts[r[COL["状態"] - 1]] = counts.get(r[COL["状態"] - 1], 0) + 1
    print(counts, "削除ログ", wb["削除ログ"].max_row - 1)


def main():
    cmd, *args = sys.argv[1:]
    if cmd == "add-sources":
        add_sources(args[0])
    elif cmd == "add-types":
        add_types(args[0])
    elif cmd == "add-candidates":
        add_candidates(args[0])
    elif cmd == "present":
        present(args[0], args[1:])
    elif cmd == "adopt":
        adopt(args[0], [int(a) for a in args[1:]])
    elif cmd == "status":
        status()
    else:
        sys.exit(__doc__)


if __name__ == "__main__":
    main()
